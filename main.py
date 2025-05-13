import os
import argparse
from openpyxl import load_workbook
from weasyprint import HTML
import markdown
import re

def escape_markdown(text):
    if not text:
        return ""
    return str(text).replace("|", "\\|").replace("\n", " ").strip()

def md_to_pdf(md_path, pdf_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        md_content = f.read()
    
    # Конвертируем Markdown в HTML
    html_content = markdown.markdown(md_content, extensions=['tables'])
    
    # Добавляем CSS стили для таблиц
    html = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="utf-8">
        <style>
            body {{
                font-family: Arial, sans-serif;
                line-height: 1.1;
                margin: 15px;
                padding-top: 15px;
            }}
            table {{
                width: 100%;
                border-collapse: collapse;
                margin: 15px 0;
                page-break-inside: avoid;
            }}
            th, td {{
                border: 1px solid #ddd;
                padding: 8px;
                text-align: left;
            }}
            th {{
                background-color: #f2f2f2;
                font-weight: bold;
            }}
            h1 {{
                color: #333;
                text-align: center;
                font-size: 22px;
            }}
            @page {{
                size: A4;
                margin: 1cm;
            }}
        </style>
    </head>
    <body>
        {html_content}
    </body>
    </html>
    """
    
    HTML(string=html, encoding='utf-8').write_pdf(pdf_path)

def get_responsible_for_mat_care(group):

  if (group == None) or (group == "—"):
    return "\n"

  text = "По вопросам матпомощи обращайся к "
  
  if (group[0] == 'M') or (group[5] == '4') or (group[5] == '3') or \
     (group == "Б05-111") or (group == "Б05-112") or (group == "Б05-120") or \
     (group == "Б05-211") or (group == "Б05-311") or (group == "Б05-411") or \
     (group == "Б05-312") or (group == "Б05-412") or (group == "Б05-225") or \
     (group == "Б05-312"):
      return text + "[Никита Вейда](https://vk.com/id238860017)\n"

  # Cкрыты другие имена из соображений конфиденциальности
  if (re.fullmatch(r"Б05-\d0\d", group)) or (re.fullmatch(r"Б05-\d5\d", group)) or (group[0] == 'А'):
    return text

  return text


def main(excel_file_path):
    wb = load_workbook(excel_file_path, data_only=True)
    
    # Получаем нужные листы
    sheet_students = wb["СТУДЕНТЫ"]
    sheet_merged = wb["MERGED"]
    
    # Считываем данные с листа "СТУДЕНТЫ"
    students_data = []
    for row in sheet_students.iter_rows(min_row=2, values_only=True):
        fio = row[0]  # ФИО
        group = row[1]  # Группа
        email = row[5]  # Почта
        
        if fio and group and email:
            students_data.append({
                "fio": fio.strip(),
                "group": group.strip(),
                "email": email.strip()
            })
    
    # Считываем все данные с листа "Merged"
    merged_data = []
    for row in sheet_merged.iter_rows(min_row=2, values_only=True):
        merged_data.append(row)
    
    # Для каждого студента ищем совпадения по ФИО и Группе в merged_data
    for student in students_data:
        fio = student["fio"]
        group = student["group"]
        email = student["email"]
        
        # Собираем строки из merged, где совпадают fio и group
        matched_rows = []
        for row in merged_data:
            merged_fio = str(row[1]).strip() if row[1] else ""
            merged_group = str(row[2]).strip() if row[2] else ""
            
            if merged_fio == fio and merged_group == group:
                matched_rows.append(row)
        
        if matched_rows:
            # Создаем директории для MD и PDF
            md_dir = os.path.join("done_markdown", email)
            pdf_dir = os.path.join("done_pdf", email)
            
            os.makedirs(md_dir, exist_ok=True)
            os.makedirs(pdf_dir, exist_ok=True)
            
            # Создаём Markdown файл
            md_content = []
            
            # Заголовок
            md_content.append("# Информация о выплатах материальной помощи\n")
            md_content.append("Mарт 2025 год\n")

            # Представитель
            md_content.append(get_responsible_for_mat_care(group))
            
            # Информация о студенте
            md_content.append(f"**{fio} {group}**\n\n")
            
            # Таблица
            headers = [
                "Номер заявления", "Категория", "Подкатегория",
                "Сумма по чекам, руб", "Назначено, руб", "Комментарий"
            ]
            
            # Заголовки таблицы
            md_content.append("| " + " | ".join(headers) + " |")
            
            # Разделительная строка
            md_content.append("|" + "|".join(["---"] * len(headers)) + "|")
            
            # Данные таблицы
            for row in matched_rows:
                row_data = [
                    escape_markdown(row[0]) if row[0] else "-",
                    escape_markdown(row[3]) if row[3] else "-",
                    escape_markdown(row[4]) if row[4] else "-",
                    escape_markdown(row[5]) if row[5] else "-",
                    escape_markdown(row[7]) if row[7] else "-",
                    escape_markdown(row[6]) if row[6] else "-"
                ]
                md_content.append("| " + " | ".join(row_data) + " |")
            
            # Сохраняем MD файл
            md_path = os.path.join(md_dir, "main.md")
            with open(md_path, 'w', encoding='utf-8') as f:
                f.write("\n".join(md_content))
            
            # Конвертируем в PDF
            pdf_path = os.path.join(pdf_dir, "main.pdf")
            md_to_pdf(md_path, pdf_path)

    print("Готово! Файлы MD и PDF сформированы при наличии совпадений.")

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("excel_file_path", type=str, help="Excel file")
    args = parser.parse_args()
    main(args.excel_file_path)
