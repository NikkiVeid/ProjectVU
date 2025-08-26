import argparse
import openpyxl

def copy_filtered_rows_and_students(source_xlsx, sheets_list_txt, target_xlsx, prefix):
  """
  1) Считывает список нужных листов из sheets_list_txt.
  2) Создаёт новую книгу target_xlsx с двумя листами:
    - "MERGED": сюда копируются строки из нужных листов,
      где значение в первом столбце начинается с prefix.
    - "СТУДЕНТЫ": если в исходной книге есть лист "СТУДЕНТЫ",
      он полностью копируется (значения) в новый лист с тем же названием.
  """

  # Читаем требуемые названия листов из текстового файла
  with open(sheets_list_txt, 'r', encoding='utf-8') as f:
    needed_sheets = [line.strip() for line in f if line.strip()]

  # Открываем исходную книгу (только значения, без формул)
  wb_source = openpyxl.load_workbook(source_xlsx, data_only=True)

  # Создаем новую книгу
  wb_target = openpyxl.Workbook()
  # Переименовываем единственный созданный по умолчанию лист в "MERGED"
  merged_sheet = wb_target.active
  merged_sheet.title = "MERGED"

  # Счётчик строк в листе "MERGED"
  merged_row = 1

  # Проходим по листам, указанным в тексте
  for sheet_name in needed_sheets:
    if sheet_name in wb_source.sheetnames:
      source_sheet = wb_source[sheet_name]

      # Итерируемся по строкам исходного листа
      for row in source_sheet.iter_rows(values_only=True):
        first_cell_value = row[0]
        if first_cell_value is not None:
          # Преобразуем в строку (на случай, если это число)
          str_first_cell_value = str(first_cell_value)
          # Проверяем, начинается ли значение в первом столбце с префикса
          if str_first_cell_value.startswith(prefix):
            # Копируем всю строку в "MERGED"
            for col_index, cell_value in enumerate(row, start=1):
              merged_sheet.cell(row=merged_row, column=col_index).value = cell_value
            merged_row += 1

  # Копируем лист "СТУДЕНТЫ", если он существует в исходной книге
  if "СТУДЕНТЫ" in wb_source.sheetnames:
    source_students_sheet = wb_source["СТУДЕНТЫ"]
    # Создаем лист в целевой книге
    target_students_sheet = wb_target.create_sheet("СТУДЕНТЫ")

    for row_index, row in enumerate(source_students_sheet.iter_rows(values_only=True), start=1):
      for col_index, cell_value in enumerate(row, start=1):
        target_students_sheet.cell(row=row_index, column=col_index).value = cell_value

  # Сохраняем результат
  wb_target.save(target_xlsx)

if __name__ == "__main__":
  """
  Пример запуска из командной строки:
  python script.py --prefix 02 --source source.xlsx --list-file 1.txt --target result.xlsx
  где:
    --prefix        - Префикс, которым должны начинаться значения в первом столбце
    --source        - Путь к исходному XLSX файлу
    --list-file     - Путь к txt-файлу со списком листов
    --target        - Путь к выходному XLSX файлу (результат)
  """

  parser = argparse.ArgumentParser(description="Копирование и фильтрация строк из нескольких листов")
  parser.add_argument("--prefix", required=True, help="Префикс (начальные символы) для фильтрации в первом столбце")
  parser.add_argument("--source", required=True, help="Путь к исходному XLSX файлу")
  parser.add_argument("--list-file", required=True, help="Путь к txt-файлу со списком обрабатываемых листов")
  parser.add_argument("--target", required=True, help="Путь к итоговому XLSX файлу")

  args = parser.parse_args()

  copy_filtered_rows_and_students(
    source_xlsx=args.source,
    sheets_list_txt=args.list_file,
    target_xlsx=args.target,
    prefix=args.prefix
  )
